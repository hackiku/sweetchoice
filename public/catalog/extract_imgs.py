import os
from openpyxl import load_workbook

# Input and output paths
input_file = "catalog.xlsx"
output_dir = "./extracted_imgs"
os.makedirs(output_dir, exist_ok=True)

# Check if the input file exists
if not os.path.exists(input_file):
    print(f"File '{input_file}' not found in the current directory.")
    exit()

print(f"Extracting images from: {input_file}")

# Load workbook
wb = load_workbook(input_file, data_only=True)

# Loop through sheets and images
image_count = 0
for sheet_name in wb.sheetnames:
    sheet = wb[sheet_name]
    if not hasattr(sheet, '_images'):
        continue  # Skip sheets without images
    
    for img in sheet._images:
        image_count += 1
        # Save the image
        img_filename = f"image_{image_count}.png"
        img_path = os.path.join(output_dir, img_filename)
        with open(img_path, "wb") as f:
            f.write(img.ref.blob)
        print(f"Extracted: {img_path}")

print(f"Done! Extracted {image_count} image(s) to {output_dir}.")
